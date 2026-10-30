# Gera 1 planilha Excel por marca: nossos produtos x catálogo do site oficial da marca (preço de varejo em US$).
# Roda no GitHub Actions (internet aberta). Nossos produtos vêm da função conc-semanal (acao omg_marcas).
import json, re, os, sys, time, unicodedata, urllib.request
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

MARCAS = {  # marca no ERP -> site oficial (loja Shopify)
    'COSRX': ['cosrx.com'],
    'Anua': ['anua.us'],
    'SKIN1004': ['skin1004.com'],
    'Beauty of Joseon': ['beautyofjoseon.com'],
    'medicube': ['medicube.us'],
    'TIRTIR': ['tirtir.us'],
    'ROUND LAB': ['roundlab.com'],
    'celimax': ['celimax.us', 'www.celimax.com'],
    'BANILA CO': ['banilausa.com'],
    'numbuzin': ['numbuzinus.com'],
    'PURITO SEOUL': ['purito.com'],
    'Rejuran': ['rejuranusa.com'],
    'VT Cosmetics': ['globalvt-cosmetics.com'],
}
NOMES_ERP = {'VT Cosmetics': ['VT Cosmetics', 'VT']}
CAMBIO, MARKUP = 5.40, 0.20
UA = {'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/129 Safari/537.36', 'Accept': 'application/json'}
FUNC = 'https://nzsiuniabovvptuqmkmi.supabase.co/functions/v1/conc-semanal'
OUT = 'tools/precos/planilhas'
os.makedirs(OUT, exist_ok=True)


def get(u, tries=3):
    for i in range(tries):
        try:
            return urllib.request.urlopen(urllib.request.Request(u, headers=UA), timeout=40).read()
        except Exception as e:
            err = e; time.sleep(3 * (i + 1))
    raise err


def post(u, body):
    r = urllib.request.Request(u, data=json.dumps(body).encode(), headers={'content-type': 'application/json'}, method='POST')
    return json.loads(urllib.request.urlopen(r, timeout=120).read())


def norm(s):
    return unicodedata.normalize('NFD', str(s or '').lower()).encode('ascii', 'ignore').decode()


STOP = set('com para de da do das dos and the with a o e em no na ml g un unidade unidades caixa x kit set pack of in for by new edition version ver'.split())
GEN = set('facial rosto pele korean skincare skin care k beauty best seller exclusive deal amazon s 1 no'.split())
TIPOS = {'serum': 'serum', 'ampoule': 'ampola', 'ampule': 'ampola', 'essence': 'essencia', 'cream': 'creme', 'toner': 'tonico', 'pad': 'pad', 'pads': 'pad',
         'mask': 'mascara', 'masks': 'mascara', 'oil': 'oleo', 'foam': 'espuma', 'gel': 'gel', 'mist': 'bruma', 'cushion': 'cushion', 'sunscreen': 'solar',
         'sun': 'solar', 'balm': 'balm', 'lotion': 'locao', 'emulsion': 'locao', 'shampoo': 'shampoo', 'tint': 'tint', 'patch': 'patch', 'patches': 'patch',
         'cleanser': 'limpeza', 'cleansing': 'limpeza', 'stick': 'stick', 'powder': 'po', 'primer': 'primer', 'concealer': 'corretivo', 'foundation': 'base',
         'eye': 'olhos', 'lip': 'labios',
         # português (descrição do nosso nome)
         'sérum': 'serum', 'creme': 'creme', 'tonico': 'tonico', 'mascara': 'mascara', 'oleo': 'oleo', 'espuma': 'espuma', 'ampola': 'ampola', 'essencia': 'essencia',
         'adesivo': 'patch', 'adesivos': 'patch', 'protetor': 'solar', 'locao': 'locao', 'bruma': 'bruma', 'balsamo': 'balm', 'discos': 'pad', 'olhos': 'olhos', 'labial': 'labios'}


def toks(s):
    return [w for w in re.findall(r'[a-z0-9]+', norm(s)) if len(w) > 1 and w not in STOP]


def tam(s):
    t = norm(s).replace(',', '.')
    m = re.search(r'(\d+(?:\.\d+)?)\s*(ml|g|oz)\b', t)
    q = None
    if m:
        q = float(m.group(1)) * (29.57 if m.group(2) == 'oz' else 1)
    u = re.search(r'(\d+)\s*(?:un\b|unid|unidades|pcs|sheets|ea\b|pads|discos|folhas)', t) or re.search(r'caixa com (\d+)', t)
    return q, (int(u.group(1)) if u else None)


KIT = re.compile(r'\b(kit|duo|trio|set|bundle|conjunto|combo|routine)\b|\s\+\s|1\+1|\b[2-9][\s-]?(un|unid|unidades|pcs|pack|packs)\b|\bx\s?[2-9]\b|caixa com', re.I)
LIXO = re.compile(r'gift|sample|free|sachet|amostra|trial|tester|shipping|insurance|protection|e-?gift|card|subscr|subscription|refill only|bonus', re.I)


def catalogo(dominio):
    prods, p = [], 1
    while p < 20:
        j = json.loads(get(f'https://{dominio}/products.json?limit=250&page={p}'))
        L = j.get('products', [])
        prods += L
        if len(L) < 250:
            break
        p += 1
    try:
        meta = json.loads(get(f'https://{dominio}/meta.json'))
    except Exception:
        meta = {}
    itens = []
    for pr in prods:
        for v in pr.get('variants', []):
            preco = float(v.get('price') or 0)
            if preco < 1:
                continue
            vt = v.get('title') or ''
            nome = pr['title'] + ('' if vt in ('Default Title', '') else ' · ' + vt)
            itens.append({'nome': nome, 'produto': pr['title'], 'variante': '' if vt == 'Default Title' else vt, 'preco': preco,
                          'cheio': float(v['compare_at_price']) if v.get('compare_at_price') and float(v['compare_at_price']) > preco else None,
                          'sku': v.get('sku') or '', 'disp': v.get('available'), 'link': f"https://{dominio}/products/{pr['handle']}" + (f"?variant={v['id']}" if len(pr.get('variants', [])) > 1 else ''),
                          'kit': bool(KIT.search(pr['title'] + ' ' + vt)), 'lixo': bool(LIXO.search(pr['title'] + ' ' + (pr.get('product_type') or '')))})
    return itens, meta


def casa(p, cat, marca):
    en = str(p['nome']).split(':')[0]
    mb = set(toks(marca))
    A = [w for w in toks(en) if w not in mb and w not in GEN]
    if not A:
        return None, 0, 'Sem par'
    ta, ua = tam(p['nome'])
    kit_nosso = bool(KIT.search(p['nome']))
    tA = {TIPOS[w] for w in toks(p['nome']) if w in TIPOS}
    best = []
    for c in cat:
        if c['lixo']:
            continue
        B = set(w for w in toks(c['nome']) if w not in mb)
        k = sum(1 for w in A if w in B)
        s = k / len(A)
        if s < 0.5:
            continue
        tB = {TIPOS[w] for w in B if w in TIPOS}
        if tA and tB and not (tA & tB):
            continue
        tb, ub = tam(c['nome'])
        conflito = (ta and tb and abs(ta - tb) / ta > 0.06) or (ua and ub and ua != ub) or (kit_nosso != c['kit'])
        extras = len([w for w in B if w not in A and w not in GEN and w not in TIPOS and not w[0].isdigit()])
        sc = s - 0.04 * extras - (0.5 if conflito else 0) + (0.15 if ta and tb and not conflito else 0)
        best.append((sc, s, extras, conflito, bool(ta and tb), c))
    if not best:
        return None, 0, 'Sem par'
    best.sort(key=lambda x: -x[0])
    sc, s, extras, conflito, tam_ok, c = best[0]
    outros = [b[0] for b in best[1:] if b[5]['produto'] != c['produto']]
    segundo = outros[0] if outros else -9
    if conflito:
        sit = 'Conferir (tamanho/kit diferente)'
    elif s >= 0.99 and extras <= 1 and (tam_ok or len(A) >= 3) and sc - segundo >= 0.1:
        sit = 'Certeiro'
    elif s >= 0.67:
        sit = 'Provável — conferir'
    else:
        sit = 'Fraco — conferir'
    return c, sc, sit


def planilha(marca, nossos, cat, meta, dominio):
    wb = Workbook()
    ws = wb.active; ws.title = 'Comparação'
    pa = wb.create_sheet('Parâmetros')
    pa['A1'], pa['B1'] = 'Câmbio US$ → R$', CAMBIO
    pa['A2'], pa['B2'] = 'Acréscimo sobre o preço do exterior', MARKUP
    pa['B2'].number_format = '0%'
    pa['A4'] = 'Mude os valores acima e a coluna "Mercadoria sugerida" recalcula.'
    pa['A5'] = f'Site oficial: https://{dominio} · loja "{meta.get("name", "")}" · moeda {meta.get("currency", "?")} · lido em {time.strftime("%d/%m/%Y %H:%M")} UTC'
    pa['A6'] = 'Regra (Guia de precificação OMG): produto sem concorrente válido → varejo no exterior em US$ + 20%, convertido = preço da mercadoria no site. Frete e impostos entram no checkout.'
    pa.column_dimensions['A'].width = 40
    cab = ['SKU', 'Nosso produto', 'No site?', 'Preço OMG hoje (R$)', 'Situação', 'Produto oficial', 'Variante', 'Preço oficial (US$)', 'Preço cheio (US$)',
           'Disponível lá?', 'Link oficial', 'Mercadoria sugerida (R$)', 'Diferença vs hoje', 'Observação']
    ws.append(cab)
    ordem = {'Certeiro': 0, 'Provável — conferir': 1, 'Conferir (tamanho/kit diferente)': 2, 'Fraco — conferir': 3, 'Sem par': 4}
    linhas = []
    for p in nossos:
        c, sc, sit = casa(p, cat, marca)
        linhas.append((ordem[sit], not p.get('ativo'), str(p['nome']), p, c, sit))
    linhas.sort(key=lambda x: (x[0], x[1], x[2]))
    cont = {}
    for _, _, _, p, c, sit in linhas:
        cont[sit] = cont.get(sit, 0) + 1
        r = ws.max_row + 1
        ws.append([p.get('sku'), p.get('nome'), 'Sim' if p.get('ativo') else 'Não', float(p['preco']) if p.get('preco') else None, sit,
                   c['produto'] if c else None, c['variante'] if c else None, c['preco'] if c else None, c['cheio'] if c else None,
                   ('Sim' if c['disp'] else 'Esgotado') if c else None, c['link'] if c else None, None, None, None])
        if c:
            ws.cell(r, 12).value = f'=ROUND(H{r}*(1+Parâmetros!$B$2)*Parâmetros!$B$1,2)'
            ws.cell(r, 13).value = f'=IF(OR(D{r}="",L{r}=""),"",D{r}/L{r}-1)'
            ws.cell(r, 11).hyperlink = c['link']; ws.cell(r, 11).style = 'Hyperlink'
    for col, w in zip(range(1, 15), [14, 60, 8, 14, 26, 46, 18, 12, 12, 11, 40, 16, 12, 30]):
        ws.column_dimensions[get_column_letter(col)].width = w
    for r in range(2, ws.max_row + 1):
        for col in (4, 12):
            ws.cell(r, col).number_format = 'R$ #,##0.00'
        for col in (8, 9):
            ws.cell(r, col).number_format = '$#,##0.00'
        ws.cell(r, 13).number_format = '0%'
        sit = ws.cell(r, 5).value
        cor = {'Certeiro': 'E3F4E8', 'Provável — conferir': 'FFF6DA', 'Conferir (tamanho/kit diferente)': 'FDEBD9', 'Fraco — conferir': 'FDEBD9'}.get(sit)
        if cor:
            ws.cell(r, 5).fill = PatternFill('solid', fgColor=cor)
    ws.freeze_panes = 'C2'
    if ws.max_row > 1:
        t = Table(displayName='Comparacao', ref=f'A1:N{ws.max_row}'); t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True); ws.add_table(t)
    # catálogo oficial completo
    cs = wb.create_sheet('Catálogo oficial')
    cs.append(['Produto oficial', 'Variante', 'Preço (US$)', 'Preço cheio (US$)', 'Disponível?', 'Kit/combo?', 'SKU deles', 'Link'])
    for c in sorted(cat, key=lambda x: x['nome']):
        if c['lixo']:
            continue
        cs.append([c['produto'], c['variante'], c['preco'], c['cheio'], 'Sim' if c['disp'] else 'Esgotado', 'Sim' if c['kit'] else '', c['sku'], c['link']])
        cs.cell(cs.max_row, 8).hyperlink = c['link']
    for col, w in zip(range(1, 9), [60, 22, 12, 14, 11, 10, 18, 50]):
        cs.column_dimensions[get_column_letter(col)].width = w
    cs.freeze_panes = 'A2'
    for row in (ws[1], cs[1]):
        for cell in row:
            cell.font = Font(bold=True)
    wb.move_sheet('Parâmetros', offset=1)
    nome = re.sub(r'[^A-Za-z0-9]+', '-', marca).strip('-')
    arq = f'{OUT}/precos-oficiais-{nome}.xlsx'
    wb.save(arq)
    return arq, cont


def main():
    todos = sum(NOMES_ERP.get(m, [m]) for m in MARCAS) if False else [n for m in MARCAS for n in NOMES_ERP.get(m, [m])]
    nossos = post(FUNC, {'acao': 'omg_marcas', 'marcas': todos, 'loja': 'x'})['produtos']
    print('nossos produtos:', len(nossos), flush=True)
    resumo = []
    for marca, doms in MARCAS.items():
        ns = [p for p in nossos if p['marca'] in NOMES_ERP.get(marca, [marca])]
        cat, meta, dom, erro = [], {}, doms[0], ''
        for d in doms:
            try:
                c, m = catalogo(d)
                if len(c) > len(cat):
                    cat, meta, dom = c, m, d
            except Exception as e:
                erro = str(e)[:120]
        if not cat:
            resumo.append({'marca': marca, 'site': dom, 'erro': erro or 'sem catálogo', 'nossos': len(ns)}); print(marca, 'ERRO', erro, flush=True); continue
        arq, cont = planilha(marca, ns, cat, meta, dom)
        resumo.append({'marca': marca, 'site': dom, 'moeda': meta.get('currency'), 'itens_oficiais': len(cat), 'nossos': len(ns), 'arquivo': arq, **cont})
        print(marca, dom, meta.get('currency'), len(cat), len(ns), cont, flush=True)
    json.dump(resumo, open(f'{OUT}/resumo.json', 'w'), ensure_ascii=False, indent=1)


if __name__ == '__main__':
    main()
